from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile

from catalog.auid import STRUCTURE_FAMILY_VALUES
from catalog.batch_upload import (
    PHASE_STATUS_VALUES,
    RAW_DATA_TYPE_VALUES,
    normalize_synthesis_steps,
)
from catalog.documents import (
    Material,
    Recipe,
    compute_material_auid,
    compute_recipe_auid,
    normalize_elements_payload,
)
from catalog.raw_db import record_raw_file


PRIMARY_XRD_PRIORITY = [".asc", ".csv", ".txt", ".xrdml", ".raw"]
SUPPORTED_EXTENSIONS = {".asc", ".csv", ".txt", ".xrdml", ".raw", ".svg"}

# Primary-file extension -> raw_data_type enum value.
EXTENSION_RAW_DATA_TYPE = {
    ".asc": "xrd",
    ".csv": "xrd",
    ".txt": "xrd",
    ".xrdml": "xrd",
    ".raw": "xrd",
}

# Manifest columns; missing columns are tolerated by the parser.
MANIFEST_HEADERS = [
    "Batch ID",
    "Material / catalyst",
    "Target composition",
    "Synthesis route",
    "XRD",
    "XRD file",
    "Reference",
    "Structure Family",
    "Phase status",
    "Spacegroup",
    "Raw data type",
    "Milling time (h)",
    "Milling rpm",
    "Atmosphere",
    "Temp profile",
    "Cooling method",
    "Notes",
]

SUBSCRIPT_MAP = str.maketrans(
    {
        "₀": "0",
        "₁": "1",
        "₂": "2",
        "₃": "3",
        "₄": "4",
        "₅": "5",
        "₆": "6",
        "₇": "7",
        "₈": "8",
        "₉": "9",
        "．": ".",
    }
)


@dataclass
class ManifestRow:
    batch_id: str
    material_name: str = ""
    target_composition: str = ""
    synthesis_route: str = ""
    xrd_status: str = ""
    # Explicit filename linking a manifest row to its XRD file, for exports
    # whose files are NOT named by batch ID (so folder/filename matching fails).
    xrd_filename: str = ""
    reference: str = ""
    structure_family: str = ""
    # Optional structured fields (expanded manifest schema).
    phase_status: str = ""
    spacegroup: str = ""
    raw_data_type: str = ""
    milling_time_hours: str = ""
    milling_rpm: str = ""
    atmosphere: str = ""
    temp_profile: str = ""
    cooling_method: str = ""
    notes: str = ""


@dataclass
class FolderFile:
    folder: str
    filename: str
    zip_path: str
    extension: str
    size_bytes: int


@dataclass
class BatchItem:
    batch_id: str
    manifest: ManifestRow | None = None
    files: list[FolderFile] = field(default_factory=list)
    primary_file: FolderFile | None = None
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


@dataclass
class BatchPreview:
    items: list[BatchItem]
    total_rows: int
    total_folders: int
    valid_count: int
    warning_count: int
    error_count: int
    zip_only: bool = False


def normalize_batch_id(value: str) -> str:
    raw = str(value or "").strip()
    match = re.search(r"SP15M[-_]?0*([0-9]+)", raw, flags=re.I)
    if match:
        return f"SP15M-{int(match.group(1)):03d}"
    return raw


XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0"


def _xlsx_to_csv_text(data: bytes, key_header: str) -> str:
    """Flatten the relevant sheet of an .xlsx workbook into CSV text.

    Picks the first worksheet whose header area contains ``key_header`` (e.g.
    "Batch ID") so a multi-sheet workbook with an "Instructions" tab still works;
    falls back to the first sheet.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)

    chosen = None
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=15, values_only=True):
            if any(str(cell).strip() == key_header for cell in row if cell is not None):
                chosen = sheet
                break
        if chosen is not None:
            break
    if chosen is None:
        chosen = workbook.worksheets[0]

    output = io.StringIO()
    writer = csv.writer(output)
    for row in chosen.iter_rows(values_only=True):
        writer.writerow(["" if cell is None else cell for cell in row])
    return output.getvalue()


def load_manifest_text(file_obj, key_header: str) -> str:
    """Return manifest CSV text from an uploaded CSV or .xlsx file.

    Excel uploads are converted internally so the rest of the pipeline only ever
    sees CSV text. ``key_header`` selects the right worksheet for .xlsx inputs.
    """
    raw = file_obj.read()
    if isinstance(raw, str):
        return raw
    if raw[:4] == XLSX_MAGIC:
        return _xlsx_to_csv_text(raw, key_header)
    if raw[:4] == XLS_MAGIC:
        raise ValueError(
            "Legacy .xls files are not supported. Save as .xlsx or CSV and re-upload."
        )
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def parse_manifest_csv(file_obj) -> dict[str, ManifestRow]:
    text = load_manifest_text(file_obj, "Batch ID")

    lines = [line for line in text.splitlines() if line.strip()]

    if not lines:
        return {}

    header_index = 0
    for index, line in enumerate(lines[:15]):
        if "Batch ID" in line:
            header_index = index
            break

    text = "\n".join(lines[header_index:])

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel

    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows: dict[str, ManifestRow] = {}

    # ManifestRow field -> accepted column headers, first non-empty wins.
    aliases = {
        "material_name": ("Material / catalyst", "material_name"),
        "target_composition": ("Target composition", "composition"),
        "synthesis_route": (
            "Synthesis route",
            "Synthesis protocol",
            "Protocol",
            "Procedure",
            "Synthesis",
            "synthesis_route",
        ),
        "xrd_status": ("XRD", "xrd_status"),
        "xrd_filename": (
            "XRD file",
            "XRD filename",
            "XRD file name",
            "File name",
            "Filename",
            "File",
            "Data file",
        ),
        "reference": ("Reference", "DOI", "reference"),
        "structure_family": ("Structure Family", "structure_family"),
        "phase_status": ("Phase status", "phase_status"),
        "spacegroup": ("Spacegroup", "spacegroup"),
        "raw_data_type": ("Raw data type", "raw_data_type"),
        "milling_time_hours": ("Milling time (h)", "milling_time_hours"),
        "milling_rpm": ("Milling rpm", "milling_rpm"),
        "atmosphere": ("Atmosphere", "atmosphere"),
        "temp_profile": ("Temp profile", "temp_profile"),
        "cooling_method": ("Cooling method", "cooling_method"),
        "notes": ("Notes", "notes"),
    }

    for row in reader:
        normalized_row = {
            str(key or "").strip(): str(value or "").strip()
            for key, value in row.items()
        }

        # Case-insensitive header lookup so "Target Composition", "TARGET
        # COMPOSITION" and "target composition" all resolve. Spreadsheets rarely
        # match the canonical casing exactly; matching case-sensitively silently
        # dropped whole columns (e.g. composition) while batch IDs still parsed.
        lower_row: dict[str, str] = {}
        for key, value in normalized_row.items():
            lower_key = key.lower()
            if value and not lower_row.get(lower_key):
                lower_row[lower_key] = value
            else:
                lower_row.setdefault(lower_key, value)

        def pick(*names):
            return next(
                (lower_row[n.lower()] for n in names if lower_row.get(n.lower())),
                "",
            )

        batch_id = normalize_batch_id(
            pick("Batch ID", "batch_id", "Sample ID", "sample_id")
        )
        if not batch_id:
            continue

        rows[batch_id] = ManifestRow(
            batch_id=batch_id,
            **{field_name: pick(*names) for field_name, names in aliases.items()},
        )

    return rows


def scan_xrd_zip(zip_file) -> dict[str, list[FolderFile]]:
    folders: dict[str, list[FolderFile]] = {}

    with zipfile.ZipFile(zip_file) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue

            parts = Path(info.filename).parts

            # Skip macOS archive cruft (__MACOSX/ resource forks, ._AppleDouble
            # files) so they don't masquerade as real XRD files.
            filename = parts[-1]
            if "__MACOSX" in parts or filename.startswith("._"):
                continue

            extension = Path(filename).suffix.lower()
            if extension not in SUPPORTED_EXTENSIONS:
                continue

            # Prefer the batch ID from a containing folder (SP15M-001/scan.txt),
            # but fall back to the batch ID embedded in the filename itself
            # (SP15M-001.txt with no per-batch subfolder — the common export
            # where each scan is named by its batch). Without the fallback these
            # flat files were dropped entirely ("No files").
            folder_part = None
            for part in parts[:-1]:
                if re.search(r"SP15M[-_]?0*[0-9]+", part, flags=re.I):
                    folder_part = part
                    break
            if folder_part is None and re.search(
                r"SP15M[-_]?0*[0-9]+", filename, flags=re.I
            ):
                folder_part = filename

            if folder_part is None:
                continue

            folder = normalize_batch_id(folder_part)

            folders.setdefault(folder, []).append(
                FolderFile(
                    folder=folder,
                    filename=filename,
                    zip_path=info.filename,
                    extension=extension,
                    size_bytes=info.file_size,
                )
            )

    return folders


def scan_all_supported_files(zip_file) -> dict[str, FolderFile]:
    """Index every supported file in the archive by lowercased basename.

    Lets a manifest link a row to its XRD by an explicit filename (the ``XRD
    file`` column) even when the file itself does not carry the batch ID — the
    case where folder/filename matching finds "No files". First occurrence of a
    given basename wins.
    """
    index: dict[str, FolderFile] = {}
    with zipfile.ZipFile(zip_file) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            parts = Path(info.filename).parts
            filename = parts[-1]
            if "__MACOSX" in parts or filename.startswith("._"):
                continue
            extension = Path(filename).suffix.lower()
            if extension not in SUPPORTED_EXTENSIONS:
                continue
            folder = normalize_batch_id(parts[-2]) if len(parts) >= 2 else ""
            index.setdefault(
                filename.lower(),
                FolderFile(
                    folder=folder,
                    filename=filename,
                    zip_path=info.filename,
                    extension=extension,
                    size_bytes=info.file_size,
                ),
            )
    return index


def _match_explicit_file(
    row: ManifestRow | None, file_index: dict[str, FolderFile] | None
) -> FolderFile | None:
    """Resolve the file a manifest row names via its ``XRD file`` column."""
    if not row or not file_index or not row.xrd_filename:
        return None
    name = row.xrd_filename.strip().lower()
    if name in file_index:
        return file_index[name]
    # Tolerate a path being given instead of a bare filename.
    base = name.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    return file_index.get(base)


def choose_primary_file(files: list[FolderFile]) -> FolderFile | None:
    for extension in PRIMARY_XRD_PRIORITY:
        candidates = [file for file in files if file.extension == extension]
        if candidates:
            return sorted(candidates, key=lambda file: file.filename)[0]
    return None


def parse_composition_formula(formula: str) -> dict[str, float]:
    if not formula:
        raise ValueError("Missing target composition")

    text = formula.translate(SUBSCRIPT_MAP)
    text = text.replace("(", "").replace(")", "")

    matches = re.findall(r"([A-Z][a-z]?)([0-9]*\.?[0-9]*)", text)
    elements: dict[str, float] = {}

    for symbol, amount in matches:
        value = float(amount) if amount else 1.0
        elements[symbol] = elements.get(symbol, 0.0) + value

    if not elements:
        raise ValueError(f"Could not parse composition: {formula}")

    return normalize_elements_payload(elements)


_TEMP_C_RE = re.compile(r"(\d+(?:\.\d+)?)\s*°?\s*[cC]\b")
_TEMP_HOLD_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:h|hr|hour)", re.I)


def build_synthesis_steps(row: ManifestRow) -> tuple[list[dict[str, Any]], list[str]]:
    """Build typed synthesis steps from the (expanded) manifest row.

    The structured columns map to typed steps (``ball_milling``,
    ``heat_treatment``, ``cooling``); the free-form ``Synthesis route`` becomes
    an ``other`` step. Raw step dicts are passed through
    :func:`catalog.batch_upload.normalize_synthesis_steps` so the stored shape
    matches the manual add form. Returns ``(steps, errors)``.
    """
    raw_steps: list[dict[str, Any]] = []

    milling: dict[str, Any] = {"step_type": "ball_milling"}
    if row.milling_time_hours:
        milling["milling_time_hours"] = row.milling_time_hours
    if row.milling_rpm:
        milling["milling_rpm"] = row.milling_rpm
    if row.atmosphere:
        milling["atmosphere"] = row.atmosphere
    if len(milling) > 1:
        raw_steps.append(milling)

    if row.temp_profile:
        heat: dict[str, Any] = {"step_type": "heat_treatment", "notes": row.temp_profile}
        temp_match = _TEMP_C_RE.search(row.temp_profile)
        if temp_match:
            heat["max_temp_c"] = temp_match.group(1)
        hold_match = _TEMP_HOLD_RE.search(row.temp_profile)
        if hold_match:
            heat["hold_time_hours"] = hold_match.group(1)
        if row.atmosphere:
            heat["atmosphere"] = row.atmosphere
        raw_steps.append(heat)

    if row.cooling_method:
        raw_steps.append({"step_type": "cooling", "cooling_method": row.cooling_method})

    if row.synthesis_route:
        raw_steps.append(
            {
                "step_type": "other",
                "description": row.synthesis_route,
                "notes": f"Batch imported from {row.batch_id}",
            }
        )

    return normalize_synthesis_steps(raw_steps)


def resolve_phase_status(row: ManifestRow) -> str:
    """Resolve a phase-status enum from an explicit column, else the XRD text."""
    explicit = (row.phase_status or "").strip().lower().replace(" ", "_")
    if explicit in PHASE_STATUS_VALUES:
        return explicit
    text = f"{row.phase_status} {row.xrd_status}".lower()
    if "multi" in text:
        return "multi_phase"
    if "single" in text:
        return "single_phase"
    return "not_confirmed"


def resolve_raw_data_type(row: ManifestRow, primary_file: "FolderFile | None") -> str | None:
    """Resolve ``raw_data_type`` from an explicit column, else the file extension."""
    explicit = (row.raw_data_type or "").strip().lower()
    if explicit in RAW_DATA_TYPE_VALUES and explicit not in ("", "na"):
        return explicit
    if primary_file is not None:
        return EXTENSION_RAW_DATA_TYPE.get(primary_file.extension, "xrd")
    return None


def resolve_spacegroup(row: ManifestRow) -> str:
    return (row.spacegroup or "").strip() or "unknown"


def validate_item(item: BatchItem, *, zip_only: bool = False) -> None:
    row = item.manifest

    if zip_only:
        if not item.files:
            item.warnings.append("No supported files found.")
        if item.files and not item.primary_file:
            item.warnings.append("No primary XRD file could be selected.")
        return

    if row is None:
        item.errors.append("Missing manifest row.")
        return

    if not row.target_composition:
        item.errors.append("Missing target composition.")
    else:
        try:
            parse_composition_formula(row.target_composition)
        except ValueError as exc:
            item.errors.append(str(exc))

    _, step_errors = build_synthesis_steps(row)
    item.errors.extend(step_errors)

    if row.structure_family and resolve_structure_family(row, "unknown") not in STRUCTURE_FAMILY_VALUES:
        item.errors.append(f"Invalid Structure Family: {row.structure_family!r}")

    explicit_phase = (row.phase_status or "").strip().lower().replace(" ", "_")
    if explicit_phase and explicit_phase not in PHASE_STATUS_VALUES:
        item.warnings.append(
            f"Unrecognized Phase status {row.phase_status!r}; will infer from XRD text."
        )

    explicit_type = (row.raw_data_type or "").strip().lower()
    if explicit_type and explicit_type not in RAW_DATA_TYPE_VALUES:
        item.warnings.append(
            f"Unrecognized Raw data type {row.raw_data_type!r}; will infer from file."
        )

    if not row.synthesis_route and not any(
        (row.milling_time_hours, row.milling_rpm, row.temp_profile, row.cooling_method)
    ):
        item.warnings.append("Missing synthesis route.")

    if not item.files:
        item.warnings.append("No XRD/supporting files found.")

    if item.files and not item.primary_file:
        item.warnings.append("No primary XRD file could be selected.")

    raw_count = len([file for file in item.files if file.extension == ".raw"])
    if raw_count > 1:
        item.warnings.append(f"Multiple .raw files found: {raw_count}")

    if "insufficient" in (row.xrd_status or "").lower() and item.files:
        item.warnings.append("Manifest says insufficient XRD sample, but files were found.")


def build_preview(
    *,
    manifest_rows: dict[str, ManifestRow],
    folders: dict[str, list[FolderFile]],
    file_index: dict[str, FolderFile] | None = None,
    zip_only: bool = False,
) -> BatchPreview:
    all_ids = sorted(set(manifest_rows.keys()) | set(folders.keys()))
    items: list[BatchItem] = []

    for batch_id in all_ids:
        files = list(folders.get(batch_id, []))
        row = manifest_rows.get(batch_id)

        # A manifest row may name its XRD explicitly (XRD file column); attach
        # that file even when it wasn't matched by batch ID, and treat it as the
        # primary so an explicit choice always wins over auto-selection.
        explicit = _match_explicit_file(row, file_index)
        if explicit is not None and all(f.zip_path != explicit.zip_path for f in files):
            files.append(explicit)

        item = BatchItem(
            batch_id=batch_id,
            manifest=row,
            files=files,
            primary_file=explicit or choose_primary_file(files),
        )
        validate_item(item, zip_only=zip_only)
        items.append(item)

    return BatchPreview(
        items=items,
        total_rows=len(manifest_rows),
        total_folders=len(folders),
        valid_count=len([item for item in items if not item.errors]),
        warning_count=len([item for item in items if item.warnings]),
        error_count=len([item for item in items if item.errors]),
        zip_only=zip_only,
    )


def create_preview(*, manifest_file=None, archive_file, zip_only: bool = False) -> BatchPreview:
    manifest_rows: dict[str, ManifestRow] = {}

    if manifest_file is not None:
        manifest_rows = parse_manifest_csv(manifest_file)

    archive_file.seek(0)
    folders = scan_xrd_zip(archive_file)
    archive_file.seek(0)
    file_index = scan_all_supported_files(archive_file)

    return build_preview(
        manifest_rows=manifest_rows,
        folders=folders,
        file_index=file_index,
        zip_only=zip_only,
    )


def generate_manifest_template_csv(preview: BatchPreview) -> str:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=MANIFEST_HEADERS)
    writer.writeheader()

    for item in preview.items:
        writer.writerow(
            {
                "Batch ID": item.batch_id,
                "XRD": "XRD completed." if item.primary_file else "",
            }
        )

    return output.getvalue()


def persist_uploaded_file(uploaded_file, prefix: str) -> str:
    suffix = Path(uploaded_file.name).suffix
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S%f")
    safe_name = uploaded_file.name.replace("/", "_").replace("\\", "_")
    path = f"{prefix}/{timestamp}_{safe_name}"
    return default_storage.save(path, uploaded_file)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def save_zip_member(zf: zipfile.ZipFile, file_info: FolderFile, batch_id: str) -> tuple[str, str, int]:
    data = zf.read(file_info.zip_path)
    file_hash = sha256_bytes(data)

    safe_filename = file_info.filename.replace("/", "_").replace("\\", "_")
    stored_path = f"experimental_batch/{batch_id}/{file_hash[:12]}_{safe_filename}"

    default_storage.save(stored_path, ContentFile(data))

    return stored_path, file_hash, len(data)


def resolve_structure_family(row: ManifestRow, default_structure_family: str) -> str:
    return (row.structure_family or default_structure_family or "unknown").strip().lower()


def commit_batch(
    *,
    manifest_path: str,
    archive_path: str,
    user,
    request,
    structure_family: str,
) -> dict[str, Any]:
    """Import a batch by routing every row through ``persist_experimental_trial``,
    sharing dedup/hashing/archiving with the manual form. Returns counters plus
    a ``skipped_details`` list."""
    # Lazy imports: views.py imports this module, so top-level would be circular.
    from catalog.synthesis_worker import enqueue_synthesis_job
    from catalog.views import DuplicateFileError, DuplicateRecordError, persist_experimental_trial

    with default_storage.open(manifest_path, "rb") as manifest_file:
        manifest_rows = parse_manifest_csv(manifest_file)

    with default_storage.open(archive_path, "rb") as archive_file:
        folders = scan_xrd_zip(archive_file)

    with default_storage.open(archive_path, "rb") as archive_file:
        file_index = scan_all_supported_files(archive_file)

    preview = build_preview(
        manifest_rows=manifest_rows,
        folders=folders,
        file_index=file_index,
        zip_only=False,
    )

    created_materials = 0
    created_recipes = 0
    created_trials = 0
    recorded_files = 0
    skipped = 0
    skipped_details: list[str] = []

    username = user.get_username()

    with default_storage.open(archive_path, "rb") as archive_file:
        with zipfile.ZipFile(archive_file) as zf:
            for item in preview.items:
                if item.errors or item.manifest is None:
                    skipped += 1
                    skipped_details.append(
                        f"{item.batch_id}: {'; '.join(item.errors) or 'no manifest row'}"
                    )
                    continue

                row = item.manifest
                row_structure_family = resolve_structure_family(row, structure_family)
                elements = parse_composition_formula(row.target_composition)
                synthesis_steps, _ = build_synthesis_steps(row)

                material_auid = compute_material_auid(elements, row_structure_family)
                recipe_auid = compute_recipe_auid(material_auid, synthesis_steps)
                material_existed = Material.objects(id=material_auid).first() is not None
                recipe_existed = Recipe.objects(id=recipe_auid).first() is not None

                # Wrap the primary ZIP member as an uploaded file so the shared
                # persist path can hash it, dedup, parse/plot, and record it.
                primary_upload = None
                if item.primary_file is not None:
                    data = zf.read(item.primary_file.zip_path)
                    primary_upload = SimpleUploadedFile(
                        item.primary_file.filename,
                        data,
                        content_type="application/octet-stream",
                    )

                try:
                    result = persist_experimental_trial(
                        user=user,
                        request=request,
                        raw_elements=elements,
                        structure_family=row_structure_family,
                        synthesis_steps=synthesis_steps,
                        phase_status=resolve_phase_status(row),
                        spacegroup=resolve_spacegroup(row),
                        element_sites={},
                        raw_data_type=resolve_raw_data_type(row, item.primary_file),
                        notes=(row.notes or row.xrd_status or None),
                        csv_file=primary_upload,
                        reject_duplicates=True,
                        source_batch_id=row.batch_id,
                    )
                except (DuplicateFileError, DuplicateRecordError) as exc:
                    skipped += 1
                    skipped_details.append(f"{row.batch_id}: {exc}")
                    continue

                if not material_existed:
                    created_materials += 1
                if not recipe_existed:
                    created_recipes += 1
                created_trials += 1
                if primary_upload is not None:
                    recorded_files += 1

                # Queue background AI discretization of the free-text route (no-op
                # unless the LLM feature is enabled). Import stays fast; the worker
                # re-keys the recipe later. See catalog/synthesis_worker.py.
                enqueue_synthesis_job(
                    kind="experiment",
                    recipe_auid=result["recipe_auid"],
                    material_auid=result["material_auid"],
                    route_text=row.synthesis_route,
                    username=username,
                    trial_id=result["trial_id"],
                )

                # The shared path only stores the single primary file; record the
                # remaining folder files so multi-file folders aren't dropped.
                for file_info in item.files:
                    if (
                        item.primary_file is not None
                        and file_info.zip_path == item.primary_file.zip_path
                    ):
                        continue
                    stored_path, file_hash, size_bytes = save_zip_member(
                        zf, file_info, item.batch_id
                    )
                    record_raw_file(
                        file_hash=file_hash,
                        material_auid=result["material_auid"],
                        recipe_auid=result["recipe_auid"],
                        trial_id=result["trial_id"],
                        original_filename=file_info.filename,
                        stored_path=stored_path,
                        content_type=file_info.extension.lstrip("."),
                        size_bytes=size_bytes,
                        uploaded_by=username,
                        elements=elements,
                        structure_family=row_structure_family,
                        notes=row.xrd_status,
                        tags=["batch_upload", "experimental", file_info.extension.lstrip(".")],
                    )
                    recorded_files += 1

    return {
        "created_materials": created_materials,
        "created_recipes": created_recipes,
        "created_trials": created_trials,
        "recorded_files": recorded_files,
        "skipped": skipped,
        "skipped_details": skipped_details,
    }